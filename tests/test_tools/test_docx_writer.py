from docx import Document

from src.tools.base_tool import ToolPermission
from src.tools.builtin.docx_writer import GenerateWordDocumentTool, _sanitise_filename
from src.tools.registry import default_registry


async def test_generates_a_real_readable_docx(tmp_path):
    tool = GenerateWordDocumentTool(root=tmp_path)
    result = await tool.run(
        title="V-101 Approval Note",
        sections=[
            {"heading": "Findings", "body": "Pressure exceeds the recommended limit."},
            {"heading": "Recommendation", "body": "Schedule inspection within 7 days."},
        ],
        tables=[
            {
                "heading": "Readings",
                "headers": ["Parameter", "Observed", "Limit"],
                "rows": [["pressure", "14.5", "12.0"]],
            }
        ],
    )
    assert result.ok is True
    out = result.output
    assert out["filename"].endswith(".docx")
    written = tmp_path / out["filename"]
    assert written.is_file()
    assert out["bytes"] == written.stat().st_size

    # Re-open with python-docx to prove it's a genuinely valid, readable file
    # -- not just bytes on disk.
    doc = Document(str(written))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "V-101 Approval Note" in full_text
    assert "Findings" in full_text
    assert "Schedule inspection within 7 days." in full_text
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "Parameter"
    assert doc.tables[0].rows[1].cells[0].text == "pressure"


async def test_default_filename_generated_when_omitted(tmp_path):
    tool = GenerateWordDocumentTool(root=tmp_path)
    result = await tool.run(title="Untitled Report")
    assert result.ok is True
    assert result.output["filename"].startswith("document_")
    assert result.output["filename"].endswith(".docx")


async def test_filename_sanitisation_strips_path_traversal():
    assert _sanitise_filename("../../etc/passwd") == "passwd.docx"
    assert _sanitise_filename("report.docx") == "report.docx"
    assert _sanitise_filename("weird name!@#.docx") == "weird_name_.docx"
    assert _sanitise_filename(None).endswith(".docx")
    assert _sanitise_filename("") .endswith(".docx")


async def test_generated_file_stays_confined_to_root(tmp_path):
    tool = GenerateWordDocumentTool(root=tmp_path)
    result = await tool.run(title="x", filename="../../escape.docx")
    assert result.ok is True
    # sanitised filename has no directory component left, so it lands
    # squarely inside root regardless of what was requested
    assert (tmp_path / result.output["filename"]).is_file()
    assert not any(p.name == "escape.docx" for p in tmp_path.parent.glob("*.docx"))


def test_write_permission_not_registered_by_default():
    registry = default_registry(allow_network=False, allow_write_filesystem=False)
    assert not registry.has("generate_word_document")


def test_write_permission_registers_when_enabled():
    registry = default_registry(allow_network=False, allow_write_filesystem=True)
    assert registry.has("generate_word_document")
    assert ToolPermission.WRITE_FILESYSTEM in registry.get("generate_word_document").permissions


def test_not_cacheable():
    assert GenerateWordDocumentTool.cacheable is False
